"""Source document metadata and version-scoped equipment discovery history."""
import copy
import json
import re
from pathlib import Path

from docx import Document

ALIASES = json.loads((Path(__file__).resolve().parents[1] / 'src/data/equipment_aliases.json').read_text(encoding='utf-8'))


def read_game_version(workbook):
    versions = set()
    for row in workbook['Intro'].iter_rows():
        for cell in row:
            match = re.search(r'updated\s+for\s+(?:War\s+Robots\s+)?v?(\d+\.\d+(?:\.\d+)?)', str(cell.value or ''), re.I)
            if match:
                versions.add(match[1])
    if len(versions) != 1:
        raise ValueError('DPS Intro must contain one unambiguous updated-for version')
    return versions.pop()


def read_methodology(path):
    return [paragraph.text.strip() for paragraph in Document(path).paragraphs if paragraph.text.strip()]


def item_key(name):
    name = re.sub(r'^(?:ultimate|ue)\s+', 'ue ', name.strip(), flags=re.I)
    key = re.sub(r'[^a-z0-9]', '', name.lower()).replace('fuijin', 'fujin')
    key = re.sub(r'swordunit$', 'sword', key)
    return ALIASES.get(key, key)


def equipment_catalog(tiers):
    catalog = {}
    for category, ranks in tiers.items():
        if category == 'disclaimers':
            continue
        names = catalog.setdefault(category, {})
        for rank in ranks.values():
            for item in rank['items']:
                for name in re.split(r'[,/.]', item['name']):
                    name = re.sub(r'[⬆⬇👁👥‼\ufe0f*]', '', name).strip()
                    if name:
                        names[item_key(name)] = name
    return catalog


def advance_equipment_state(version, tiers, previous=None, releases=()):
    if not re.fullmatch(r'\d+\.\d+(?:\.\d+)?', version):
        raise ValueError('Invalid game version')
    if previous and tuple(map(int, version.split('.'))) < tuple(map(int, previous['version'].split('.'))):
        raise ValueError('Equipment version moved backward; restore matching source files')
    catalog = equipment_catalog(tiers)
    history = {category: {item_key(key): value for key, value in names.items()}
               for category, names in copy.deepcopy(previous['first_seen']).items()} if previous else {}
    seed = {(category, item_key(name)): release['version']
            for release in releases if release.get('version')
            for category, names in release['items'].items() for name in names}
    new_by_category = {}
    for category, names in catalog.items():
        known = history.setdefault(category, {})
        new_by_category[category] = []
        for key, name in names.items():
            if key not in known:
                known[key] = seed.get((category, key), version if previous else None)
            if known[key] == version:
                new_by_category[category].append(name)
    return {'version': version, 'first_seen': history, 'new_by_category': new_by_category}
