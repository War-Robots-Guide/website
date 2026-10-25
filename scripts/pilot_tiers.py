"""Optional public pilot-ranking snapshot; independent of core guide ingestion."""
import argparse
from datetime import datetime, timezone
import io
from http.client import IncompleteRead
import json
import os
from pathlib import Path
import re
import sys
import tempfile
import time
import unicodedata
from urllib.error import HTTPError, URLError
from urllib.request import urlopen
import warnings
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from defusedxml.common import DefusedXmlException
from lxml.etree import XMLSyntaxError

SPREADSHEET_ID = '1RBLXWZDD-PmFxW4ChUOLNrzin8Gzxjp8tfw2gSeol8Y'
SOURCE_URL = f'https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit#gid=0'
EXPORT_URL = f'https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/export?format=xlsx'
DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / 'src/data/pilot_tiers.json'
MAX_BYTES = 2 * 1024 * 1024


def validate_pilot_tiers(data: object) -> list[str]:
    if not isinstance(data, dict):
        return ['expected an object']
    errors = []
    if type(data.get('schemaVersion')) is not int or data['schemaVersion'] != 1:
        errors.append('schemaVersion must be 1')
    if data.get('sourceUrl') != SOURCE_URL:
        errors.append('sourceUrl must identify the pilot source spreadsheet')
    stamp = data.get('fetchedAt')
    try:
        if not isinstance(stamp, str) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z', stamp):
            raise ValueError()
        datetime.strptime(stamp, '%Y-%m-%dT%H:%M:%SZ')
    except ValueError:
        errors.append('fetchedAt must be a valid UTC timestamp')
    tiers = data.get('tiers')
    if not isinstance(tiers, list) or not tiers:
        return errors + ['tiers must be a nonempty array']
    labels = set()
    for index, tier in enumerate(tiers):
        prefix = f'tiers[{index}]'
        if not isinstance(tier, dict):
            errors.append(f'{prefix} must be an object')
            continue
        label = tier.get('label')
        if not isinstance(label, str) or not label.strip():
            errors.append(f'{prefix}.label must be nonempty text')
        elif label.strip().casefold() in labels:
            errors.append(f'{prefix}.label must be unique')
        else:
            labels.add(label.strip().casefold())
        items = tier.get('items')
        if not isinstance(items, list) or not items:
            errors.append(f'{prefix}.items must be a nonempty array')
        else:
            for position, item in enumerate(items):
                if not isinstance(item, dict) or not isinstance(item.get('name'), str) or not item['name'].strip():
                    errors.append(f'{prefix}.items[{position}].name must be nonempty text')
    return errors


def parse_pilot_tiers(workbook, *, fetched_at: str) -> dict:
    headings = [(sheet, row[0].row) for sheet in workbook.worksheets
                for row in sheet.iter_rows(min_col=1, max_col=1)
                if row[0].data_type not in ('f', 'e') and isinstance(row[0].value, str)
                and row[0].value.strip().casefold() == 'pilots']
    if len(headings) != 1:
        raise ValueError('Expected exactly one Pilots heading in the workbook')
    sheet, start = headings[0]
    tiers, seen = [], set()
    for cells in sheet.iter_rows(min_row=start + 1, max_col=2):
        label, names = (cell.value for cell in cells)
        if all(value is None or (isinstance(value, str) and not value.strip()) for value in (label, names)):
            continue
        if isinstance(label, str) and label.strip().casefold() in ('maps', 'gamemodes'):
            break
        if any(cell.data_type in ('f', 'e') for cell in cells) or not all(isinstance(value, str) and value.strip() for value in (label, names)):
            raise ValueError(f'Expected literal tier label and pilot names at {sheet.title}!A{cells[0].row}:B{cells[0].row}')
        items = [{'name': name.strip()} for name in names.split(',') if name.strip()]
        for item in items:
            key = re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKD', item['name']).lower())
            if key in seen:
                warnings.warn(f'Duplicate pilot ranking preserved: {item["name"]}', stacklevel=2)
            seen.add(key)
        tiers.append({'label': label.strip(), 'items': items})
    data = {'schemaVersion': 1, 'sourceUrl': SOURCE_URL, 'fetchedAt': fetched_at, 'tiers': tiers}
    errors = validate_pilot_tiers(data)
    if errors:
        raise ValueError('; '.join(errors))
    return data


def download_workbook() -> bytes:
    for attempt in range(3):
        try:
            with urlopen(EXPORT_URL, timeout=10) as response:
                content = response.read(MAX_BYTES + 1)
            if len(content) > MAX_BYTES:
                raise ValueError('Pilot spreadsheet exceeds 2 MiB download limit')
            return content
        except HTTPError as error:
            error.close()
            if error.code not in (429, 500, 502, 503, 504) or attempt == 2:
                raise
        except (URLError, TimeoutError, ConnectionError, IncompleteRead):
            if attempt == 2:
                raise
        time.sleep(2 ** attempt)
    raise RuntimeError('Unreachable download retry state')


def _parse_bytes(content):
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    except (KeyError, IndexError) as error:
        raise ValueError('Invalid pilot workbook structure') from error
    try:
        return parse_pilot_tiers(workbook, fetched_at=datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'))
    finally:
        workbook.close()


def _write_snapshot(output, data):
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode='w', encoding='utf8', dir=output.parent, delete=False) as target:
            temporary = Path(target.name)
            json.dump(data, target, ensure_ascii=False, indent=2)
            target.write('\n')
        os.replace(temporary, output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _read_valid_snapshot(path):
    data = json.loads(path.read_text(encoding='utf8'))
    errors = validate_pilot_tiers(data)
    if errors:
        raise ValueError('; '.join(errors))
    return data


def refresh_pilot_tiers(output: Path, *, fetch=download_workbook) -> bool:
    output = Path(output)
    try:
        data = _parse_bytes(fetch())
    except (OSError, ValueError, BadZipFile, InvalidFileException, ParseError, XMLSyntaxError, DefusedXmlException, IncompleteRead) as error:
        try:
            _read_valid_snapshot(output)
        except (OSError, ValueError) as fallback_error:
            raise ValueError(f'Pilot source unavailable and no valid fallback: {fallback_error}') from error
        # Escape workflow command delimiters in upstream messages.
        message = str(error).replace('%', '%25').replace('\r', '%0D').replace('\n', '%0A')
        print(f'::warning::Pilot refresh failed; keeping validated snapshot: {message}')
        return False
    _write_snapshot(output, data)
    print('Refreshed pilot tier snapshot successfully.')
    return True


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument('--refresh', action='store_true')
    mode.add_argument('--input', type=Path)
    mode.add_argument('--validate', type=Path)
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    try:
        if args.validate:
            _read_valid_snapshot(args.validate)
        elif args.input:
            _write_snapshot(args.output, _parse_bytes(args.input.read_bytes()))
        else:
            refresh_pilot_tiers(args.output, fetch=download_workbook)
        return 0
    except (OSError, ValueError, BadZipFile, InvalidFileException) as error:
        print(f'Pilot tiers: {error}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
