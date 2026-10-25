import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from content_metadata import read_game_version, read_methodology, advance_equipment_state


def catalog(*names):
    return {'Robots': {'S': {'items': [{'name': n} for n in names]}},
            'Mothership Turrets': {'A': {'items': [{'name': 'Defender System'}]}}}


class ContentMetadataTests(unittest.TestCase):
    def test_reads_version_from_intro_and_rejects_missing_marker(self):
        wb = openpyxl.Workbook()
        wb.active.title = 'Intro'
        wb.active['A3'] = 'Currently updated for War Robots v12.3'
        self.assertEqual(read_game_version(wb), '12.3')
        wb.active['A3'] = 'Introduction'
        with self.assertRaisesRegex(ValueError, 'version'):
            read_game_version(wb)

    def test_methodology_uses_document_paragraphs_not_workbook_directions(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'Methodology text.docx'
            doc = Document()
            doc.add_paragraph('What is a Value Rating?')
            doc.add_paragraph('An indicator of return on investment.')
            doc.save(path)
            self.assertEqual(read_methodology(path), ['What is a Value Rating?', 'An indicator of return on investment.'])

    def test_rollover_clears_every_category_and_accumulates_additions(self):
        initial = advance_equipment_state('12.3', catalog('Fang'))
        self.assertEqual(initial['new_by_category']['Robots'], [])
        added = advance_equipment_state('12.3', catalog('Fang', 'Dedopali'), initial)
        self.assertEqual(added['new_by_category']['Robots'], ['Dedopali'])
        rerun = advance_equipment_state('12.3', catalog('Fang', 'Dedopali', 'UE Shell'), added)
        self.assertCountEqual(rerun['new_by_category']['Robots'], ['Dedopali', 'UE Shell'])
        rollover = advance_equipment_state('12.4', catalog('Fang', 'Dedopali', 'UE Shell'), rerun)
        self.assertTrue(all(not names for names in rollover['new_by_category'].values()))
        with self.assertRaisesRegex(ValueError, 'backward'):
            advance_equipment_state('12.3', catalog('Fang'), rollover)

    def test_renames_tags_removal_and_reintroduction_do_not_make_items_new(self):
        state = advance_equipment_state('12.3', catalog('UE Sword', 'Fujin'))
        state = advance_equipment_state('12.4', catalog('⬆️Ultimate Sword Unit*', 'Fuijin'), state)
        self.assertEqual(state['new_by_category']['Robots'], [])
        state = advance_equipment_state('12.4', catalog('Fujin'), state)
        state = advance_equipment_state('12.5', catalog('UE Sword', 'Fujin'), state)
        self.assertEqual(state['new_by_category']['Robots'], [])

    def test_seed_records_are_version_scoped_not_latest_per_category(self):
        releases = [{'version': '12.3', 'items': {'Robots': ['Fang']}},
                    {'version': '12.4', 'items': {'Robots': ['Dedopali']}}]
        state = advance_equipment_state('12.3', catalog('Fang', 'Dedopali'), releases=releases)
        self.assertEqual(state['new_by_category']['Robots'], ['Fang'])
        self.assertEqual(state['new_by_category']['Mothership Turrets'], [])
        state = advance_equipment_state('12.4', catalog('Fang', 'Dedopali'), state, releases)
        self.assertEqual(state['new_by_category']['Robots'], ['Dedopali'])

    def test_family_members_have_independent_history(self):
        tiers = {'Robot Weapons': {'X': {'items': [{'name': '👁️Iaraghi L, Iaraghi M. Iaraghi H'}]}}}
        state = advance_equipment_state('12.3', tiers)
        tiers['Robot Weapons']['X']['items'].append({'name': 'New L, New M'})
        state = advance_equipment_state('12.4', tiers, state)
        self.assertEqual(state['new_by_category']['Robot Weapons'], ['New L', 'New M'])

    def test_spelling_corrections_do_not_make_existing_weapons_new(self):
        tiers = {'Robot Weapons': {'X': {'items': [{'name': 'Farenheit, Machiana'}]}}}
        state = advance_equipment_state('12.3', tiers)
        tiers['Robot Weapons']['X']['items'][0]['name'] = 'Fahrenheit, Machiara'
        state = advance_equipment_state('12.4', tiers, state)
        self.assertEqual(state['new_by_category']['Robot Weapons'], [])
