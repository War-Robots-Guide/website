"""Offline pilot-ranking ingestion and recovery contracts."""
import contextlib
import io
import json
import sys
import unittest
import warnings
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from urllib.error import HTTPError, URLError

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pilot_tiers import (SOURCE_URL, download_workbook, main, parse_pilot_tiers,
                         refresh_pilot_tiers, validate_pilot_tiers)

STAMP = '2026-10-25T04:00:00Z'


def workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Renamed'
    ws.append(['Maps'])
    ws.append(['Good maps', 'Castle'])
    ws['A40'] = ' Pilots '
    ws.append(['Hot', 'Erika Storm, Sophie Kovač'])
    ws.append([None, None])
    ws.append(['Get out', 'Jad Parkes, Yang Lee (Harpy / Siren)'])
    ws.append(['Gamemodes'])
    ws.append(['Good', '7v7'])
    return wb


def payload(wb=None):
    out = io.BytesIO()
    (wb or workbook()).save(out)
    return out.getvalue()


def snapshot():
    return {'schemaVersion': 1, 'sourceUrl': SOURCE_URL, 'fetchedAt': STAMP,
            'tiers': [{'label': 'Hot', 'items': [{'name': 'Erika Storm'}]}]}


class PilotTierTests(unittest.TestCase):
    def test_parse_moved_pilots_preserves_source_text(self):
        data = parse_pilot_tiers(workbook(), fetched_at=STAMP)
        self.assertEqual(data['tiers'], [
            {'label': 'Hot', 'items': [{'name': 'Erika Storm'}, {'name': 'Sophie Kovač'}]},
            {'label': 'Get out', 'items': [{'name': 'Jad Parkes'}, {'name': 'Yang Lee (Harpy / Siren)'}]},
        ])
        self.assertEqual(validate_pilot_tiers(data), [])

    def test_missing_ambiguous_and_empty_sections_reject_update(self):
        for variant in ['missing', 'duplicate', 'empty', 'no names', 'formula', 'numeric', 'error']:
            with self.subTest(variant=variant):
                wb = workbook()
                if variant == 'missing': wb.active['A40'] = 'Other'
                elif variant == 'duplicate': wb.create_sheet()['A1'] = 'pilots'
                elif variant == 'empty': wb.active.delete_rows(41, 3)
                elif variant == 'no names': wb.active['B41'] = None
                elif variant == 'formula': wb.active['B41'] = '=1+1'
                elif variant == 'numeric': wb.active['B41'] = 123
                else: wb.active['B41'] = '#N/A'
                with self.assertRaises(ValueError):
                    parse_pilot_tiers(wb, fetched_at=STAMP)

    def test_duplicates_are_preserved_with_warning(self):
        wb = workbook()
        wb.active['B43'] = 'Erika Storm, Sati "Cat" Felidae, Kevin_O_Reilly'
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            data = parse_pilot_tiers(wb, fetched_at=STAMP)
        self.assertIn('Erika Storm', str(caught[0].message))
        self.assertEqual(data['tiers'][1]['items'], [{'name': 'Erika Storm'}, {'name': 'Sati "Cat" Felidae'}, {'name': 'Kevin_O_Reilly'}])

    def test_schema_rejects_invalid_contracts_but_allows_duplicate_names(self):
        for key, value in [('schemaVersion', True), ('schemaVersion', 2), ('sourceUrl', 'https://example.com'),
                           ('fetchedAt', '2026-02-30T00:00:00Z'), ('fetchedAt', 'bad'), ('tiers', []),
                           ('tiers', [{'label': '', 'items': []}]), ('tiers', [{'label': 'Hot', 'items': [{'name': ''}]}])]:
            with self.subTest(key=key, value=value):
                data = snapshot(); data[key] = value
                self.assertTrue(validate_pilot_tiers(data))
        self.assertTrue(validate_pilot_tiers(None))
        self.assertTrue(validate_pilot_tiers({}))
        data = snapshot(); data['tiers'].append(data['tiers'][0])
        self.assertTrue(validate_pilot_tiers(data))
        data['tiers'][1] = {'label': 'No', 'items': [{'name': 'Erika Storm'}]}
        self.assertEqual(validate_pilot_tiers(data), [])

    def test_refresh_replaces_invalid_fallback_with_valid_download(self):
        with TemporaryDirectory() as folder:
            output = Path(folder) / 'tiers.json'; output.write_text('broken')
            self.assertTrue(refresh_pilot_tiers(output, fetch=payload))
            data = json.loads(output.read_text())
            self.assertEqual(validate_pilot_tiers(data), [])
            self.assertEqual(data['tiers'][0]['label'], 'Hot')

    def test_source_failures_leave_valid_fallback_bytes_and_timestamp_unchanged(self):
        for failure in [b'<html>login</html>', b'PKbroken', URLError('offline'), TimeoutError('timeout')]:
            with self.subTest(failure=failure), TemporaryDirectory() as folder:
                output = Path(folder) / 'tiers.json'
                original = json.dumps(snapshot()).encode(); output.write_bytes(original)
                def fetch():
                    if isinstance(failure, Exception): raise failure
                    return failure
                with contextlib.redirect_stdout(io.StringIO()) as messages:
                    self.assertFalse(refresh_pilot_tiers(output, fetch=fetch))
                self.assertEqual(output.read_bytes(), original)
                self.assertIn('::warning::', messages.getvalue())

    def test_no_usable_fallback_and_output_write_failures_are_fatal(self):
        with TemporaryDirectory() as folder:
            output = Path(folder) / 'tiers.json'
            with self.assertRaises(ValueError): refresh_pilot_tiers(output, fetch=lambda: b'bad')
            output.write_text(json.dumps(snapshot())); original = output.read_bytes()
            with patch('pilot_tiers.os.replace', side_effect=PermissionError('denied')):
                with self.assertRaises(PermissionError): refresh_pilot_tiers(output, fetch=payload)
            self.assertEqual(output.read_bytes(), original)
            self.assertEqual(list(Path(folder).iterdir()), [output])

    def test_download_retries_transient_failures_and_limits_size(self):
        response = io.BytesIO(b'xlsx')
        with patch('pilot_tiers.urlopen', side_effect=[URLError('offline'), HTTPError('url', 503, 'busy', {}, None), response]) as request, patch('pilot_tiers.time.sleep'):
            self.assertEqual(download_workbook(), b'xlsx')
            self.assertEqual(request.call_count, 3)
            self.assertEqual(request.call_args.kwargs['timeout'], 10)
        for code in [403, 404]:
            with patch('pilot_tiers.urlopen', side_effect=HTTPError('url', code, 'no', {}, None)) as request:
                with self.assertRaises(HTTPError): download_workbook()
                self.assertEqual(request.call_count, 1)
        with patch('pilot_tiers.urlopen', return_value=io.BytesIO(b'x' * (2 * 1024 * 1024 + 1))):
            with self.assertRaises(ValueError): download_workbook()
        with patch('pilot_tiers.urlopen', side_effect=TimeoutError('timeout')) as request, patch('pilot_tiers.time.sleep'):
            with self.assertRaises(TimeoutError): download_workbook()
            self.assertEqual(request.call_count, 3)

    def test_cli_refresh_exit_status_distinguishes_fallback_from_total_failure(self):
        with TemporaryDirectory() as folder, patch('pilot_tiers.download_workbook', side_effect=URLError('offline')):
            output = Path(folder) / 'tiers.json'
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(main(['--refresh', '--output', str(output)]), 1)
                output.write_text(json.dumps(snapshot()))
                self.assertEqual(main(['--refresh', '--output', str(output)]), 0)
                self.assertEqual(main(['--validate', str(output)]), 0)


if __name__ == '__main__':
    unittest.main()
